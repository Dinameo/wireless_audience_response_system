import re
import unicodedata
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class ImportFormatError(ValueError):
    pass


QUESTION_RE = re.compile(
    r"^(?:(?:câu|question)[ \t]*)?(\d+)\s*[.):-]?\s*(.*)$", re.IGNORECASE
)
OPTION_RE = re.compile(r"^([ABCD])\s*[.):-]\s*(.+)$", re.IGNORECASE)
ANSWER_RE = re.compile(
    r"^(?:đáp\s*án(?:\s*đúng)?|dap\s*an(?:\s*dung)?|answer|correct\s*answer)"
    r"\s*[:=-]?\s*([ABCD])(?:\b|$)",
    re.IGNORECASE,
)
POINTS_RE = re.compile(
    r"^(?:điểm|diem|points?)\s*[:=-]?\s*[\d.,]+\s*$", re.IGNORECASE
)
ANSWER_KEY_RE = re.compile(r"(\d+)\s*[.):-]?\s*[-=]?\s*([ABCD])\b", re.IGNORECASE)


def _plain(value):
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = text.replace("đ", "d")
    return " ".join(
        "".join(char for char in text if unicodedata.category(char) != "Mn")
        .replace("_", " ")
        .replace("-", " ")
        .split()
    )


def _document_text(content: bytes, suffix: str) -> str:
    try:
        if suffix == ".docx":
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ImportFormatError(
            f"Không thể đọc file {suffix}. File bị lỗi hoặc không đúng định dạng; "
            "vui lòng lưu lại tài liệu và thử tải lên lần nữa."
        ) from exc
    raise ImportFormatError("Chỉ hỗ trợ file .docx hoặc .pdf")


def parse_questions(content: bytes, filename: str) -> list[dict]:
    suffix = Path(filename or "").suffix.lower()
    text = _document_text(content, suffix)
    lines = [" ".join(line.split()) for line in text.splitlines() if line.strip()]
    if not lines:
        detail = "PDF ảnh scan không có lớp chữ" if suffix == ".pdf" else "file không có nội dung"
        raise ImportFormatError(f"Không đọc được câu hỏi: {detail}.")

    answer_key = {}
    answer_key_lines = set()
    for index, line in enumerate(lines):
        normalized = _plain(line)
        if normalized.startswith(("dap an", "answer key")):
            pairs = ANSWER_KEY_RE.findall(line)
            if pairs:
                answer_key.update({int(number): choice.upper() for number, choice in pairs})
                answer_key_lines.add(index)

    blocks = []
    current = None
    current_field = None

    def finish():
        nonlocal current
        if current:
            blocks.append(current)
            current = None

    for index, line in enumerate(lines):
        if index in answer_key_lines:
            continue
        question_match = QUESTION_RE.match(line)
        option_match = OPTION_RE.match(line)
        answer_match = ANSWER_RE.match(line)
        points_match = POINTS_RE.match(line)

        # Một dòng chỉ là số không được xem là câu hỏi mới.
        if question_match and (question_match.group(2) or _plain(line).startswith(("cau", "question"))):
            finish()
            current = {
                "number": int(question_match.group(1)),
                "text_parts": [question_match.group(2)] if question_match.group(2) else [],
                "options": {},
                "correct_choice": None,
            }
            current_field = "text"
        elif current and option_match:
            choice = option_match.group(1).upper()
            current["options"][choice] = option_match.group(2).strip()
            current_field = choice
        elif current and answer_match:
            current["correct_choice"] = answer_match.group(1).upper()
            current_field = None
        elif current and points_match:
            # Exams always split 10 points evenly; old per-question point lines are ignored.
            current_field = None
        elif current:
            if current_field in ("A", "B", "C", "D"):
                current["options"][current_field] += " " + line
            elif current_field == "text":
                current["text_parts"].append(line)
    finish()

    if not blocks:
        raise ImportFormatError(
            "Không nhận diện được câu hỏi. Hãy dùng dạng “Câu 1: ...”, “A. ...” và “Đáp án: A”."
        )

    questions = []
    errors = []
    seen_numbers = set()
    for block in blocks:
        number = block["number"]
        text_value = " ".join(block["text_parts"]).strip()
        correct = block["correct_choice"] or answer_key.get(number)
        missing = [choice for choice in "ABCD" if not block["options"].get(choice)]
        if number in seen_numbers:
            errors.append(f"Câu {number}: số câu bị trùng")
        if not text_value:
            errors.append(f"Câu {number}: thiếu nội dung")
        if missing:
            errors.append(f"Câu {number}: thiếu lựa chọn {', '.join(missing)}")
        if not correct:
            errors.append(f"Câu {number}: thiếu đáp án đúng")
        seen_numbers.add(number)
        if text_value and not missing and correct:
            questions.append({
                "text": text_value[:1000],
                "a": block["options"]["A"][:500],
                "b": block["options"]["B"][:500],
                "c": block["options"]["C"][:500],
                "d": block["options"]["D"][:500],
                "correct_choice": correct,
            })

    if errors:
        raise ImportFormatError("; ".join(errors))
    return questions


HEADER_ALIASES = {
    "mssv": {"mssv", "ma sinh vien", "student id", "student code"},
    "name": {"ho ten", "ho va ten", "ten sinh vien", "name", "full name"},
    "device_id": {"ma thiet bi", "ma nut", "device", "device id", "button id"},
}


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_students(content: bytes, filename: str) -> list[dict]:
    if Path(filename or "").suffix.lower() != ".xlsx":
        raise ImportFormatError("Danh sách sinh viên phải là file Excel .xlsx")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
    except Exception as exc:
        raise ImportFormatError(f"Không thể đọc file Excel: {exc}") from exc

    header_row = None
    columns = {}
    rows = list(sheet.iter_rows(values_only=True))
    for row_index, row in enumerate(rows[:20], start=1):
        normalized = [_plain(value) for value in row]
        found = {}
        for field, aliases in HEADER_ALIASES.items():
            for column_index, title in enumerate(normalized):
                if title in aliases:
                    found[field] = column_index
                    break
        if 'mssv' in found and 'name' in found:
            header_row, columns = row_index, found
            break
    if not header_row:
        raise ImportFormatError(
            "Không tìm thấy hai cột bắt buộc: MSSV và Họ tên."
        )

    students = []
    errors = []
    seen_mssv = set()
    for excel_row, row in enumerate(rows[header_row:], start=header_row + 1):
        values = {
            field: _cell_text(row[index] if index < len(row) else None)
            for field, index in columns.items()
        }
        values.setdefault('device_id', '')
        if not any(values.values()):
            continue
        missing = [field for field in ('mssv', 'name') if not values[field]]
        if missing:
            errors.append(f"Dòng {excel_row}: có ô bắt buộc đang trống")
            continue
        if len(values["mssv"]) > 30 or len(values["name"]) > 120 or len(values["device_id"]) > 40:
            errors.append(f"Dòng {excel_row}: dữ liệu dài quá giới hạn")
            continue
        if values["mssv"] in seen_mssv:
            errors.append(f"Dòng {excel_row}: MSSV {values['mssv']} bị trùng trong file")
            continue
        seen_mssv.add(values["mssv"])
        students.append(values)

    if errors:
        raise ImportFormatError("; ".join(errors))
    if not students:
        raise ImportFormatError("File Excel không có dòng sinh viên nào.")
    return students
