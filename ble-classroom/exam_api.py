"""Class–subject-scoped exams, rosters, answers, and gateway device presence."""
import hashlib
import re
import secrets
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Literal

import pymysql
from fastapi import APIRouter, File, Form, Header, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, ValidationError, field_validator

from importers import ImportFormatError, parse_questions, parse_students
from workspace import class_idle, course_idle, own_class, own_course, own_exam, own_question, purge_exam, purge_question, teacher_id

MAX_UPLOAD_SIZE = 10 * 1024 * 1024
MAX_EXAM_POINTS_MICRO = 10_000_000
SCORE_PRECISION = Decimal('0.01')
RETAKE_SUFFIX = re.compile(r'(?:\s*·\s*Thi lại(?:\s*\(\d+\))?)+\s*$', re.IGNORECASE)


def retake_root(title):
    """Remove prior retake suffixes before creating the next attempt."""
    match = RETAKE_SUFFIX.search(title)
    root = title[:match.start()].strip() if match else title.strip()
    return root or 'Bài thi', len(re.findall(r'·\s*Thi lại', match.group() if match else '', re.IGNORECASE))


def exam_title_taken(db, course_id, title, exclude_id=None):
    sql = 'SELECT id FROM exams WHERE course_id=%s AND title=%s'
    values = [course_id, title]
    if exclude_id is not None:
        sql += ' AND id<>%s'
        values.append(exclude_id)
    db.execute(sql + ' LIMIT 1', values)
    return bool(db.fetchone())


def unique_retake_title(db, course_id, source_title, exclude_id=None):
    root, _ = retake_root(source_title)
    attempt = 1
    while True:
        suffix = ' · Thi lại' if attempt == 1 else f' · Thi lại ({attempt})'
        candidate = root[:200 - len(suffix)].rstrip() + suffix
        if not exam_title_taken(db, course_id, candidate, exclude_id):
            return candidate
        attempt += 1


def unique_legacy_title(db, course_id, title, exclude_id):
    """Give older duplicate names a stable suffix before adding the unique key."""
    if RETAKE_SUFFIX.search(title):
        return unique_retake_title(db, course_id, title, exclude_id)
    attempt = 2
    while True:
        suffix = f' ({attempt})'
        candidate = title[:200 - len(suffix)].rstrip() + suffix
        if not exam_title_taken(db, course_id, candidate, exclude_id):
            return candidate
        attempt += 1


def init_exam_title_rules(database):
    """Repair old repeated retake labels and enforce course-scoped unique names."""
    with database() as db:
        db.execute('''SELECT 1 FROM information_schema.STATISTICS
            WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='exams'
            AND INDEX_NAME='uq_exam_course_title' LIMIT 1''')
        if db.fetchone():
            return
        db.execute('SELECT id,course_id,title FROM exams WHERE course_id IS NOT NULL ORDER BY id FOR UPDATE')
        for exam in db.fetchall():
            _, repeats = retake_root(exam['title'])
            if repeats > 1:
                title = unique_retake_title(db, exam['course_id'], exam['title'], exam['id'])
                db.execute('UPDATE exams SET title=%s,revision=revision+1 WHERE id=%s', (title, exam['id']))
        db.execute('''SELECT course_id,title FROM exams WHERE course_id IS NOT NULL
            GROUP BY course_id,title HAVING COUNT(*)>1''')
        for duplicate in db.fetchall():
            db.execute('''SELECT id,title FROM exams WHERE course_id=%s AND title=%s ORDER BY id FOR UPDATE''',
                       (duplicate['course_id'], duplicate['title']))
            for exam in db.fetchall()[1:]:
                title = unique_legacy_title(db, duplicate['course_id'], exam['title'], exam['id'])
                db.execute('UPDATE exams SET title=%s,revision=revision+1 WHERE id=%s', (title, exam['id']))
    with database() as db:
        db.execute('ALTER TABLE exams ADD UNIQUE KEY uq_exam_course_title (course_id,title)')


def equal_exam_points(count):
    """Split exactly 10 points, with at most 0.000001 between questions."""
    if count < 1:
        raise ValueError('Bài thi cần ít nhất một câu hỏi')
    base, remainder = divmod(MAX_EXAM_POINTS_MICRO, count)
    return [Decimal(base + (position < remainder)) / 1_000_000 for position in range(count)]


def rounded_score(value):
    return float(Decimal(value).quantize(SCORE_PRECISION, rounding=ROUND_HALF_UP))


def init_score_scale(database):
    """Upgrade integer point columns and normalize earlier exams once."""
    with database() as db:
        db.execute('''SELECT DATA_TYPE,NUMERIC_SCALE FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='questions' AND COLUMN_NAME='points' ''')
        column = db.fetchone()
        if column['DATA_TYPE'] != 'decimal' or column['NUMERIC_SCALE'] != 6:
            db.execute('ALTER TABLE questions MODIFY COLUMN points DECIMAL(10,6) NOT NULL')
    with database() as db:
        db.execute("SELECT version FROM workspace_migrations WHERE version='exam-score-10-v1' FOR UPDATE")
        if db.fetchone():
            return
        db.execute('SELECT id FROM exams ORDER BY id')
        for exam in db.fetchall():
            db.execute('SELECT id FROM questions WHERE exam_id=%s ORDER BY position,id', (exam['id'],))
            question_ids = [row['id'] for row in db.fetchall()]
            if question_ids:
                db.executemany('UPDATE questions SET points=%s WHERE id=%s',
                               list(zip(equal_exam_points(len(question_ids)), question_ids)))
        db.execute('UPDATE questions SET points=10 WHERE exam_id IS NULL AND points>10')
        db.execute("INSERT INTO workspace_migrations(version) VALUES('exam-score-10-v1')")


class StudentIn(BaseModel):
    mssv: str = Field(min_length=1, max_length=30)
    name: str = Field(min_length=1, max_length=120)
    device_id: str = Field(default='', max_length=40)
    preferred_protocol: Literal['BLE', 'WIFI'] = 'BLE'

    @field_validator('mssv', 'name')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Không được để trống')
        return value

    @field_validator('device_id')
    @classmethod
    def trim_device(cls, value):
        return value.strip()


class StudentEdit(StudentIn):
    revision: int = Field(ge=0)


class QuestionIn(BaseModel):
    text: str = Field(min_length=1, max_length=1000)
    a: str = Field(min_length=1, max_length=500)
    b: str = Field(min_length=1, max_length=500)
    c: str = Field(min_length=1, max_length=500)
    d: str = Field(min_length=1, max_length=500)
    correct_choice: str = Field(pattern='^[ABCD]$')
    points: Decimal = Field(default=Decimal('10'), gt=0, le=10, decimal_places=6)

    @field_validator('text', 'a', 'b', 'c', 'd')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Không được để trống')
        return value


class ExamTitle(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    revision: int = Field(ge=0)

    @field_validator('title')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Tên bài thi không được trống')
        return value


class ExamEdit(ExamTitle):
    questions: list[QuestionIn] = Field(min_length=1, max_length=500)
    time_limit_seconds: int | None = Field(default=None, ge=5, le=3600)


class DeviceAssignment(BaseModel):
    student_id: int = Field(gt=0)
    device_id: str = Field(min_length=1, max_length=40)
    protocol: Literal['BLE', 'WIFI'] | None = None

    @field_validator('device_id')
    @classmethod
    def trim_id(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Mã thiết bị không được trống')
        return value


class StartExamIn(BaseModel):
    revision: int = Field(ge=0)
    assignments: list[DeviceAssignment] = Field(min_length=1)


class AnswerIn(BaseModel):
    question_id: int = Field(gt=0)
    device_id: str = Field(min_length=1, max_length=40)
    choice: str = Field(pattern='^[ABCD]$')
    sequence: int = Field(ge=0)


class PresenceIn(BaseModel):
    question_id: int = Field(gt=0)
    device_id: str = Field(min_length=1, max_length=40)
    protocol: Literal['BLE', 'USB', 'WIFI', 'SIMULATED'] = 'BLE'
    gateway_id: str = Field(min_length=1, max_length=64)
    connected: bool
    rssi: int | None = Field(default=None, ge=-127, le=0)

    @field_validator('device_id', 'gateway_id')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Không được để trống')
        return value


class BleOnlineIn(BaseModel):
    device_id: str = Field(min_length=1, max_length=40)
    gateway_id: str = Field(min_length=1, max_length=64)
    connected: bool
    gatt_profile: Literal['split', 'pdf'] | None = None
    authorized: bool = False

    @field_validator('device_id', 'gateway_id')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Không được để trống')
        return value


class BleAuthorizeIn(BaseModel):
    device_id: str = Field(min_length=1, max_length=40)

    @field_validator('device_id')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Không được để trống')
        return value


class WifiAnswerIn(BaseModel):
    question_id: int = Field(gt=0)
    choice: str = Field(pattern='^[ABCD]$')
    sequence: int = Field(ge=0)


class WifiPresenceIn(BaseModel):
    question_id: int = Field(gt=0)
    connected: bool
    rssi: int | None = Field(default=None, ge=-127, le=0)
    gateway_id: str = Field(default='esp32-wifi', min_length=1, max_length=64)

    @field_validator('gateway_id')
    @classmethod
    def trim_gateway(cls, value):
        value = value.strip()
        if not value:
            raise ValueError('Không được để trống')
        return value


async def upload_bytes(file):
    content = await file.read(MAX_UPLOAD_SIZE + 1)
    if not content:
        raise HTTPException(422, 'File rỗng')
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(413, 'File vượt quá giới hạn 10 MB')
    return content


def enroll(db, item, course_id, owner):
    db.execute('SELECT class_id FROM class_subjects WHERE id=%s', (course_id,))
    course = db.fetchone()
    if not course:
        raise HTTPException(404, 'Không tìm thấy lớp–môn')
    return enroll_class(db, item, course['class_id'], owner)


def enroll_class(db, item, class_id, owner):
    """Attach one teacher-owned student to a class and all its subjects."""
    db.execute('SELECT * FROM students WHERE teacher_id=%s AND mssv=%s FOR UPDATE',
               (owner, item['mssv']))
    student = db.fetchone()
    if student:
        if student['name'] != item['name'] or (item['device_id'] and student['device_id'] != item['device_id']):
            raise HTTPException(409, 'MSSV đã có hồ sơ khác. Hãy sửa hồ sơ trước khi thêm vào lớp.')
        ident = student['id']
    else:
        try:
            db.execute('INSERT INTO students(teacher_id,mssv,name,device_id,preferred_protocol) VALUES(%s,%s,%s,%s,%s)',
                       (owner, item['mssv'], item['name'], item['device_id'], item.get('preferred_protocol', 'BLE')))
            ident = db.lastrowid
        except pymysql.err.IntegrityError:
            raise HTTPException(409, 'MSSV đã được dùng trong tài khoản này.')
    if item['device_id']:
        db.execute('''INSERT IGNORE INTO devices(teacher_id,device_id) VALUES(%s,%s)''',
                   (owner, item['device_id']))
    db.execute('SELECT active FROM class_students WHERE class_id=%s AND student_id=%s',
               (class_id, ident))
    member = db.fetchone()
    if member and member['active']:
        raise HTTPException(409, 'Sinh viên đã có trong lớp này')
    db.execute('''INSERT INTO class_students(class_id,student_id,active) VALUES(%s,%s,TRUE)
        ON DUPLICATE KEY UPDATE active=TRUE''', (class_id, ident))
    db.execute('''INSERT INTO enrollments(course_id,student_id,active)
        SELECT id,%s,TRUE FROM class_subjects WHERE class_id=%s
        ON DUPLICATE KEY UPDATE active=TRUE''', (ident, class_id))
    return ident


def roster(db, exam=None, question=None):
    if exam and exam['status'] != 'draft':
        db.execute('SELECT student_id AS id,mssv,name,device_id FROM exam_students WHERE exam_id=%s ORDER BY mssv', (exam['id'],))
    elif question and question['status'] != 'draft':
        db.execute('SELECT student_id AS id,mssv,name,device_id FROM question_students WHERE question_id=%s ORDER BY mssv', (question['id'],))
    else:
        course_id = exam['course_id'] if exam else question['course_id']
        db.execute('''SELECT s.id,s.mssv,s.name,s.device_id,s.preferred_protocol FROM enrollments en JOIN students s ON s.id=en.student_id
            WHERE en.course_id=%s AND en.active=TRUE ORDER BY s.mssv''', (course_id,))
    return db.fetchall()


def ble_authorization(db, device_id):
    """An ID is accepted before an exam; an open question additionally gates by roster."""
    db.execute('''SELECT q.id,q.exam_id,c.teacher_id FROM questions q
        JOIN class_subjects cs ON cs.id=q.course_id JOIN classes c ON c.id=cs.class_id
        WHERE q.status='open' LIMIT 1''')
    question = db.fetchone()
    if question:
        db.execute("SELECT 1 FROM devices WHERE teacher_id=%s AND device_id=%s AND protocol='BLE' LIMIT 1",
                   (question['teacher_id'], device_id))
    else:
        db.execute("SELECT 1 FROM devices WHERE device_id=%s AND protocol='BLE' LIMIT 1",
                   (device_id,))
    if not db.fetchone():
        return {'authorized': False, 'reason': 'unknown_device'}
    if question:
        if question['exam_id'] is not None:
            db.execute('SELECT 1 FROM exam_students WHERE exam_id=%s AND device_id=%s LIMIT 1',
                       (question['exam_id'], device_id))
        else:
            db.execute('SELECT 1 FROM question_students WHERE question_id=%s AND device_id=%s LIMIT 1',
                       (question['id'], device_id))
        if not db.fetchone():
            return {'authorized': False, 'reason': 'not_in_roster'}
    return {'authorized': True, 'reason': 'known_device'}


def validate_exam_assignments(people, assignments):
    """Every participant gets exactly one device; a device serves one person per attempt."""
    expected = {person['id'] for person in people}
    mapping = {}
    devices = set()
    for assignment in assignments:
        if assignment.student_id not in expected or assignment.student_id in mapping:
            raise HTTPException(422, 'Danh sách gán thiết bị không khớp danh sách sinh viên của bài thi.')
        device_key = assignment.device_id.casefold()
        if device_key in devices:
            raise HTTPException(422, f'Mã {assignment.device_id} đang được gán cho hai sinh viên trong cùng bài thi.')
        mapping[assignment.student_id] = assignment.device_id
        devices.add(device_key)
    if set(mapping) != expected:
        raise HTTPException(422, 'Hãy chọn đúng một thiết bị cho mỗi sinh viên trước khi bắt đầu.')
    return mapping


def exam_summary_data(db, exam):
    people = roster(db, exam=exam)
    db.execute('SELECT id,position,text,correct_choice,points FROM questions WHERE exam_id=%s ORDER BY position', (exam['id'],))
    questions = db.fetchall()
    db.execute('SELECT a.student_id,a.question_id,a.choice FROM answers a JOIN questions q ON q.id=a.question_id WHERE q.exam_id=%s', (exam['id'],))
    answers = {(a['student_id'],a['question_id']):a['choice'] for a in db.fetchall()}
    for person in people:
        score = Decimal('0')
        person.update(score=0, answers=[])
        for q in questions:
            choice = answers.get((person['id'],q['id']))
            correct = choice is not None and choice == q['correct_choice']
            awarded = q['points'] if correct else Decimal('0')
            score += awarded
            person['answers'].append({'position':q['position'], 'choice':choice,
                'correct_choice':q['correct_choice'], 'correct':correct,
                'points':rounded_score(awarded), 'points_exact':float(awarded)})
        person['score'] = rounded_score(score)
    return {'exam':exam, 'questions':questions, 'students':people,
            'max_score':rounded_score(sum((q['points'] for q in questions), Decimal('0')))}


def export_score_workbook(summary):
    """Export score and every answer as numeric Excel data, never spreadsheet formulas."""
    book = Workbook()
    scores = book.active
    scores.title = 'Bảng điểm'
    detail = book.create_sheet('Chi tiết câu trả lời')
    scores.append(['MSSV', 'Họ tên', 'Mã thiết bị đã dùng', 'Số câu đúng', 'Số câu sai',
                   'Chưa trả lời', 'Điểm', 'Điểm tối đa'])
    detail.append(['MSSV', 'Họ tên', 'Mã thiết bị đã dùng', 'Câu', 'Đã chọn',
                   'Đáp án đúng', 'Kết quả', 'Điểm đạt'])

    def safe_text(value):
        value = str(value or '')
        return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value

    for person in summary['students']:
        replies = person['answers']
        correct = sum(answer['correct'] for answer in replies)
        missing = sum(answer['choice'] is None for answer in replies)
        scores.append([safe_text(person['mssv']), safe_text(person['name']),
                       safe_text(person['device_id']), correct, len(replies)-correct-missing,
                       missing, person['score'], summary['max_score']])
        for answer in replies:
            detail.append([safe_text(person['mssv']), safe_text(person['name']),
                           safe_text(person['device_id']), answer['position'],
                           answer['choice'] or '', answer['correct_choice'],
                           'Đúng' if answer['correct'] else 'Chưa trả lời' if answer['choice'] is None else 'Sai',
                           answer.get('points_exact', answer['points'])])
    for sheet, widths in [(scores, [18, 28, 24, 15, 15, 16, 12, 14]),
                          (detail, [18, 28, 24, 9, 12, 15, 17, 14])]:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64+column)].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='08745F')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if sheet is scores and cell.column in (7, 8):
                    cell.number_format = '0.00'
                elif sheet is detail and cell.column == 8:
                    cell.number_format = '0.000000'
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def insert_questions(db, items, exam_id, course_id):
    points = equal_exam_points(len(items))
    db.executemany('''INSERT INTO questions(text,a,b,c,d,correct_choice,points,exam_id,course_id,position)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        [(q['text'],q['a'],q['b'],q['c'],q['d'],q['correct_choice'],points[p-1],exam_id,course_id,p)
         for p,q in enumerate(items,1)])


def advance_exam(db, exam_id, question_id):
    db.execute("UPDATE questions SET status='closed' WHERE id=%s", (question_id,))
    db.execute("SELECT id,position FROM questions WHERE exam_id=%s AND status='draft' ORDER BY position LIMIT 1 FOR UPDATE", (exam_id,))
    following = db.fetchone()
    if following:
        db.execute("UPDATE questions SET status='open' WHERE id=%s", (following['id'],))
        db.execute('''UPDATE exams SET current_position=%s,revision=revision+1,
            deadline_at=CASE WHEN time_limit_seconds IS NULL THEN NULL
                ELSE DATE_ADD(UTC_TIMESTAMP(3),INTERVAL time_limit_seconds SECOND) END WHERE id=%s''',
            (following['position'], exam_id))
        return {'advanced':True, 'completed':False, 'position':following['position']}
    db.execute("UPDATE exams SET status='closed',current_position=NULL,deadline_at=NULL,revision=revision+1 WHERE id=%s", (exam_id,))
    return {'advanced':True, 'completed':True, 'position':None}


def reset_exam_for_restart(db, exam_id):
    """Clear the previous attempt and return the exam to a fresh draft state."""
    db.execute('SELECT id FROM questions WHERE exam_id=%s', (exam_id,))
    question_ids = [row['id'] for row in db.fetchall()]
    if question_ids:
        placeholders = ','.join(['%s'] * len(question_ids))
        db.execute(f'DELETE FROM answers WHERE question_id IN ({placeholders})', tuple(question_ids))
        db.execute(f'DELETE FROM question_students WHERE question_id IN ({placeholders})', tuple(question_ids))
    db.execute('DELETE FROM exam_students WHERE exam_id=%s', (exam_id,))
    db.execute('DELETE FROM device_presence WHERE exam_id=%s', (exam_id,))
    db.execute('DELETE FROM device_reports WHERE exam_id=%s', (exam_id,))
    db.execute("UPDATE questions SET status='draft' WHERE exam_id=%s", (exam_id,))
    db.execute("UPDATE exams SET status='draft',current_position=NULL,deadline_at=NULL,revision=revision+1 WHERE id=%s", (exam_id,))


def expire_due_exams(database):
    """Move expired questions server-side even when no browser or gateway is connected."""
    with database() as db:
        db.execute("SELECT id FROM exams WHERE status='open' AND deadline_at<=UTC_TIMESTAMP(3)")
        due_ids = [row['id'] for row in db.fetchall()]
    owners = set()
    for exam_id in due_ids:
        with database() as db:
            db.execute('''SELECT course_id,status,
                (deadline_at IS NOT NULL AND deadline_at<=UTC_TIMESTAMP(3)) AS expired
                FROM exams WHERE id=%s FOR UPDATE''', (exam_id,))
            exam = db.fetchone()
            if not exam or exam['status']!='open' or not exam['expired']:
                continue
            db.execute("SELECT id FROM questions WHERE exam_id=%s AND status='open' FOR UPDATE", (exam_id,))
            current = db.fetchone()
            if not current:
                continue
            advance_exam(db, exam_id, current['id'])
            db.execute('''SELECT c.teacher_id FROM class_subjects cs JOIN classes c ON c.id=cs.class_id
                WHERE cs.id=%s''', (exam['course_id'],))
            owners.add(db.fetchone()['teacher_id'])
    return owners


def configure_exam_api(app, database, changed, gateway_key):
    router = APIRouter(prefix='/api')

    @router.get('/classes/{class_id}/students')
    def students_in_class(class_id: int, request: Request,
                          course_id: int | None = Query(default=None, gt=0)):
        with database() as db:
            own_class(db, class_id, teacher_id(request))
            if course_id is not None:
                course = own_course(db, course_id, teacher_id(request))
                if course['class_id'] != class_id:
                    raise HTTPException(404, 'Môn không thuộc lớp đang chọn')
            db.execute('''SELECT s.id,s.mssv,s.name,s.device_id,s.preferred_protocol,s.revision,
                MAX(d.token_hash IS NOT NULL) AS wifi_enabled,
                MAX(CASE WHEN bi.connected=TRUE AND bi.last_seen>=UTC_TIMESTAMP(3)-INTERVAL 15 SECOND
                    THEN 1 ELSE 0 END) AS ble_online,
                MAX(CASE WHEN bi.connected=TRUE AND bi.authorized=TRUE AND bi.last_seen>=UTC_TIMESTAMP(3)-INTERVAL 15 SECOND
                    THEN 1 ELSE 0 END) AS ble_authorized,
                MAX(bi.gatt_profile) AS ble_profile,
                ROUND(SUM(CASE WHEN q.course_id=%s AND e.status='closed'
                    AND a.choice=q.correct_choice THEN q.points ELSE 0 END) /
                    NULLIF((SELECT COUNT(*) FROM exams completed WHERE completed.course_id=%s
                        AND completed.status='closed'),0),2) AS score
                FROM class_students cls JOIN students s ON s.id=cls.student_id
                LEFT JOIN devices d ON d.teacher_id=s.teacher_id AND d.device_id=s.device_id
                LEFT JOIN ble_device_inventory bi ON bi.device_id=s.device_id
                LEFT JOIN answers a ON a.student_id=s.id LEFT JOIN questions q ON q.id=a.question_id
                LEFT JOIN exams e ON e.id=q.exam_id
                WHERE cls.class_id=%s AND cls.active=TRUE AND s.teacher_id=%s
                GROUP BY s.id,s.mssv,s.name,s.device_id,s.preferred_protocol,s.revision ORDER BY s.mssv''',
                (course_id, course_id, class_id, teacher_id(request)))
            return db.fetchall()

    @router.post('/classes/{class_id}/students', status_code=201)
    async def create_class_student(class_id: int, item: StudentIn, request: Request):
        with database() as db:
            own_class(db, class_id, teacher_id(request), True)
            class_idle(db, class_id)
            ident = enroll_class(db, item.model_dump(), class_id, teacher_id(request))
        await changed(teacher_id(request))
        return {'id': ident, 'class_id': class_id}

    @router.post('/classes/{class_id}/import/students', status_code=201)
    async def import_class_students(class_id: int, request: Request, file: UploadFile = File(...)):
        with database() as db:
            own_class(db, class_id, teacher_id(request))
        try:
            items = parse_students(await upload_bytes(file), file.filename or '')
        except ImportFormatError as exc:
            raise HTTPException(422, str(exc))
        created, skipped = [], []
        with database() as db:
            own_class(db, class_id, teacher_id(request), True)
            class_idle(db, class_id)
            for item in items:
                try:
                    enroll_class(db, item, class_id, teacher_id(request))
                    created.append(item['mssv'])
                except HTTPException as exc:
                    skipped.append(item['mssv'] + ': ' + exc.detail)
        if created:
            await changed(teacher_id(request))
        return {'imported': len(created), 'skipped': skipped}

    @router.patch('/classes/{class_id}/students/{student_id}')
    async def edit_class_student(class_id: int, student_id: int, item: StudentEdit, request: Request):
        try:
            with database() as db:
                own_class(db, class_id, teacher_id(request), True)
                class_idle(db, class_id)
                db.execute('''SELECT s.id,s.revision,s.device_id FROM students s
                    JOIN class_students cls ON cls.student_id=s.id
                    WHERE s.id=%s AND s.teacher_id=%s AND cls.class_id=%s AND cls.active=TRUE FOR UPDATE''',
                    (student_id, teacher_id(request), class_id))
                current = db.fetchone()
                if not current:
                    raise HTTPException(404, 'Không tìm thấy sinh viên trong lớp')
                if current['revision'] != item.revision:
                    raise HTTPException(409, 'Hồ sơ đã thay đổi. Hãy tải lại trước khi sửa.')
                db.execute('''SELECT 1 FROM questions q
                    LEFT JOIN exam_students es ON es.exam_id=q.exam_id AND es.student_id=%s
                    LEFT JOIN question_students qs ON qs.question_id=q.id AND qs.student_id=%s
                    WHERE q.status='open' AND (es.student_id IS NOT NULL OR qs.student_id IS NOT NULL) LIMIT 1''',
                    (student_id, student_id))
                if db.fetchone():
                    raise HTTPException(409, 'Sinh viên đang tham gia một lượt thi. Hãy kết thúc trước khi sửa hồ sơ.')
                if item.preferred_protocol == 'BLE' and item.device_id:
                    db.execute('SELECT token_hash FROM devices WHERE teacher_id=%s AND device_id=%s',
                               (teacher_id(request), item.device_id))
                    device = db.fetchone()
                    if device and device['token_hash']:
                        raise HTTPException(409, 'Mã bo này đã kích hoạt ESP32. Chọn mã nRF khác hoặc thu hồi Wi-Fi trước.')
                db.execute('''UPDATE students SET mssv=%s,name=%s,device_id=%s,preferred_protocol=%s,revision=revision+1
                    WHERE id=%s''', (item.mssv, item.name, item.device_id, item.preferred_protocol, student_id))
                if item.device_id:
                    db.execute('INSERT IGNORE INTO devices(teacher_id,device_id) VALUES(%s,%s)',
                               (teacher_id(request), item.device_id))
        except pymysql.err.IntegrityError:
            raise HTTPException(409, 'MSSV đã được dùng trong tài khoản này.')
        await changed(teacher_id(request))
        return {'id': student_id, 'revision': item.revision + 1}

    @router.delete('/classes/{class_id}/students/{student_id}')
    async def remove_class_student(class_id: int, student_id: int, request: Request):
        with database() as db:
            own_class(db, class_id, teacher_id(request), True)
            class_idle(db, class_id)
            db.execute('''UPDATE class_students cls JOIN students s ON s.id=cls.student_id
                SET cls.active=FALSE WHERE cls.class_id=%s AND cls.student_id=%s
                AND cls.active=TRUE AND s.teacher_id=%s''',
                (class_id, student_id, teacher_id(request)))
            if not db.rowcount:
                raise HTTPException(404, 'Không tìm thấy sinh viên trong lớp')
            db.execute('''UPDATE enrollments en JOIN class_subjects cs ON cs.id=en.course_id
                SET en.active=FALSE WHERE cs.class_id=%s AND en.student_id=%s''',
                (class_id, student_id))
        await changed(teacher_id(request))
        return {'id': student_id, 'removed': True}

    @router.get('/students')
    def students(request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request))
            db.execute('''SELECT s.id,s.mssv,s.name,s.device_id,s.preferred_protocol,s.revision,
                MAX(d.token_hash IS NOT NULL) AS wifi_enabled,
                MAX(CASE WHEN bi.connected=TRUE AND bi.last_seen>=UTC_TIMESTAMP(3)-INTERVAL 15 SECOND
                    THEN 1 ELSE 0 END) AS ble_online,
                MAX(CASE WHEN bi.connected=TRUE AND bi.authorized=TRUE AND bi.last_seen>=UTC_TIMESTAMP(3)-INTERVAL 15 SECOND
                    THEN 1 ELSE 0 END) AS ble_authorized,
                MAX(bi.gatt_profile) AS ble_profile,
                ROUND(SUM(CASE WHEN q.course_id=%s AND e.status='closed'
                    AND a.choice=q.correct_choice THEN q.points ELSE 0 END) /
                    NULLIF((SELECT COUNT(*) FROM exams completed WHERE completed.course_id=%s
                        AND completed.status='closed'),0),2) AS score
                FROM enrollments en JOIN students s ON s.id=en.student_id
                LEFT JOIN devices d ON d.teacher_id=s.teacher_id AND d.device_id=s.device_id
                LEFT JOIN ble_device_inventory bi ON bi.device_id=s.device_id
                LEFT JOIN answers a ON a.student_id=s.id LEFT JOIN questions q ON q.id=a.question_id
                LEFT JOIN exams e ON e.id=q.exam_id
                WHERE en.course_id=%s AND en.active=TRUE AND s.teacher_id=%s
                GROUP BY s.id,s.mssv,s.name,s.device_id,s.preferred_protocol,s.revision ORDER BY s.mssv''',
                (course_id, course_id, course_id, teacher_id(request)))
            return db.fetchall()

    @router.post('/students', status_code=201)
    async def create_student(item: StudentIn, request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            course = own_course(db, course_id, teacher_id(request), True)
            course_idle(db, course_id)
            class_idle(db, course['class_id'])
            ident = enroll(db, item.model_dump(), course_id, teacher_id(request))
        await changed(teacher_id(request))
        return {'id':ident}

    @router.patch('/students/{student_id}')
    async def edit_student(student_id: int, item: StudentEdit, request: Request, course_id: int = Query(gt=0)):
        try:
            with database() as db:
                own_course(db, course_id, teacher_id(request), True)
                course_idle(db, course_id)
                db.execute('''SELECT s.id,s.revision,s.device_id FROM students s JOIN enrollments en ON en.student_id=s.id
                    WHERE s.id=%s AND s.teacher_id=%s AND en.course_id=%s AND en.active=TRUE FOR UPDATE''',
                    (student_id, teacher_id(request), course_id))
                current = db.fetchone()
                if not current:
                    raise HTTPException(404, 'Không tìm thấy sinh viên trong lớp–môn')
                if current['revision'] != item.revision:
                    raise HTTPException(409, 'Hồ sơ đã thay đổi. Hãy tải lại trước khi sửa.')
                db.execute("""SELECT 1 FROM questions q
                    LEFT JOIN exam_students es ON es.exam_id=q.exam_id AND es.student_id=%s
                    LEFT JOIN question_students qs ON qs.question_id=q.id AND qs.student_id=%s
                    WHERE q.status='open' AND (es.student_id IS NOT NULL OR qs.student_id IS NOT NULL) LIMIT 1""",
                    (student_id, student_id))
                if db.fetchone():
                    raise HTTPException(409, 'Sinh viên đang tham gia một lượt thi. Hãy kết thúc trước khi sửa hồ sơ.')
                if item.preferred_protocol == 'BLE' and item.device_id:
                    db.execute('SELECT token_hash FROM devices WHERE teacher_id=%s AND device_id=%s',
                               (teacher_id(request), item.device_id))
                    device = db.fetchone()
                    if device and device['token_hash']:
                        raise HTTPException(409, 'Mã bo này đã kích hoạt ESP32. Chọn mã nRF khác hoặc thu hồi Wi-Fi trước.')
                db.execute('UPDATE students SET mssv=%s,name=%s,device_id=%s,preferred_protocol=%s,revision=revision+1 WHERE id=%s',
                           (item.mssv, item.name, item.device_id, item.preferred_protocol, student_id))
                if item.device_id:
                    db.execute('INSERT IGNORE INTO devices(teacher_id,device_id) VALUES(%s,%s)',
                               (teacher_id(request), item.device_id))
        except pymysql.err.IntegrityError:
            raise HTTPException(409, 'MSSV đã được dùng trong tài khoản này')
        await changed(teacher_id(request))
        return {'id':student_id, 'revision':item.revision+1}

    @router.delete('/students/{student_id}')
    async def remove_student(student_id: int, request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            course = own_course(db, course_id, teacher_id(request), True)
            class_idle(db, course['class_id'])
            db.execute('''UPDATE class_students cls JOIN students s ON s.id=cls.student_id
                SET cls.active=FALSE WHERE cls.class_id=%s AND cls.student_id=%s
                AND cls.active=TRUE AND s.teacher_id=%s''',
                (course['class_id'], student_id, teacher_id(request)))
            if not db.rowcount:
                raise HTTPException(404, 'Không tìm thấy sinh viên trong lớp')
            db.execute('''UPDATE enrollments en JOIN class_subjects cs ON cs.id=en.course_id
                SET en.active=FALSE WHERE cs.class_id=%s AND en.student_id=%s''',
                (course['class_id'], student_id))
        await changed(teacher_id(request))
        return {'removed':True, 'history_preserved':True}

    @router.post('/students/{student_id}/wifi-token', status_code=201)
    async def issue_wifi_token(student_id: int, request: Request, course_id: int = Query(gt=0)):
        token = secrets.token_urlsafe(32)
        token_hash = hashlib.sha256(token.encode()).hexdigest()
        with database() as db:
            own_course(db, course_id, teacher_id(request), True)
            db.execute('''SELECT s.device_id,s.preferred_protocol FROM students s JOIN enrollments en ON en.student_id=s.id
                WHERE s.id=%s AND s.teacher_id=%s AND en.course_id=%s AND en.active=TRUE FOR UPDATE''',
                (student_id, teacher_id(request), course_id))
            student = db.fetchone()
            if not student:
                raise HTTPException(404, 'Không tìm thấy sinh viên trong lớp–môn')
            if not student['device_id']:
                raise HTTPException(409, 'Hãy nhập mã thiết bị ESP32 trong hồ sơ trước khi kích hoạt.')
            if student['preferred_protocol'] != 'WIFI':
                raise HTTPException(409, 'Hãy chọn ESP32 · Wi-Fi trong hồ sơ sinh viên trước khi kích hoạt.')
            db.execute('''INSERT INTO devices(teacher_id,device_id,protocol,token_hash)
                VALUES(%s,%s,'WIFI',%s) ON DUPLICATE KEY UPDATE
                protocol='WIFI',token_hash=VALUES(token_hash),created_at=CURRENT_TIMESTAMP''',
                (teacher_id(request), student['device_id'], token_hash))
        await changed(teacher_id(request))
        return {'student_id':student_id,'device_id':student['device_id'],'token':token,
                'message':'Mã gắn với thiết bị, không gắn cố định với sinh viên. Sao chép ngay; server chỉ lưu bản băm.'}

    @router.delete('/students/{student_id}/wifi-token')
    async def revoke_wifi_token(student_id: int, request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request), True)
            db.execute('''SELECT s.id,s.device_id FROM students s JOIN enrollments en ON en.student_id=s.id
                WHERE s.id=%s AND s.teacher_id=%s AND en.course_id=%s AND en.active=TRUE FOR UPDATE''',
                (student_id, teacher_id(request), course_id))
            student = db.fetchone()
            if not student:
                raise HTTPException(404, 'Không tìm thấy sinh viên trong lớp–môn')
            db.execute('''UPDATE devices SET token_hash=NULL,protocol='BLE'
                WHERE teacher_id=%s AND device_id=%s AND token_hash IS NOT NULL''',
                (teacher_id(request), student['device_id']))
            revoked = bool(db.rowcount)
        await changed(teacher_id(request))
        return {'revoked':revoked}

    @router.post('/import/students', status_code=201)
    async def import_students(request: Request, file: UploadFile = File(...), course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request))
        try:
            items = parse_students(await upload_bytes(file), file.filename or '')
        except ImportFormatError as exc:
            raise HTTPException(422, str(exc))
        created, skipped = [], []
        with database() as db:
            course = own_course(db, course_id, teacher_id(request), True)
            course_idle(db, course_id)
            class_idle(db, course['class_id'])
            for item in items:
                try:
                    enroll(db, item, course_id, teacher_id(request))
                    created.append(item['mssv'])
                except HTTPException as exc:
                    skipped.append(item['mssv']+': '+exc.detail)
        if created:
            await changed(teacher_id(request))
        return {'imported':len(created), 'skipped':skipped}

    @router.get('/exams')
    def exams(request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request))
            db.execute('''SELECT e.id,e.title,e.status,e.current_position,e.created_at,e.revision,e.time_limit_seconds,
                CASE WHEN e.deadline_at IS NULL THEN NULL ELSE
                    GREATEST(0,TIMESTAMPDIFF(MICROSECOND,UTC_TIMESTAMP(3),e.deadline_at) DIV 1000) END AS remaining_ms,
                COUNT(q.id) AS question_count,SUM(q.status='closed') AS completed_count,
                MAX(CASE WHEN q.status='open' THEN q.text END) AS current_question,
                MAX(CASE WHEN q.status='open' THEN q.a END) AS current_a,
                MAX(CASE WHEN q.status='open' THEN q.b END) AS current_b,
                MAX(CASE WHEN q.status='open' THEN q.c END) AS current_c,
                MAX(CASE WHEN q.status='open' THEN q.d END) AS current_d,
                SUM(CASE WHEN q.status='open' THEN (SELECT COUNT(*) FROM answers a WHERE a.question_id=q.id) ELSE 0 END) AS current_submitted,
                (SELECT COUNT(*) FROM exam_students es WHERE es.exam_id=e.id) AS participant_count
                FROM exams e LEFT JOIN questions q ON q.exam_id=e.id WHERE e.course_id=%s
                GROUP BY e.id,e.title,e.status,e.current_position,e.created_at,e.revision,e.time_limit_seconds,e.deadline_at
                ORDER BY e.id DESC''', (course_id,))
            return db.fetchall()

    @router.post('/import/questions', status_code=201)
    async def import_questions(request: Request, file: UploadFile = File(...), title: str = Form(default=''),
                               time_limit_seconds: int | None = Form(default=None, ge=5, le=3600),
                               course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request))
        try:
            items = parse_questions(await upload_bytes(file), file.filename or '')
            items = [QuestionIn(**item).model_dump() for item in items]
        except ImportFormatError as exc:
            raise HTTPException(422, str(exc))
        except ValidationError:
            raise HTTPException(422, 'File có câu hỏi không hợp lệ')
        name = title.strip() or Path(file.filename or 'Bài thi').stem
        if not name or len(name)>200:
            raise HTTPException(422, 'Tên bài thi phải từ 1 đến 200 ký tự')
        with database() as db:
            own_course(db, course_id, teacher_id(request), True)
            try:
                db.execute('INSERT INTO exams(title,course_id,time_limit_seconds) VALUES(%s,%s,%s)',
                           (name, course_id, time_limit_seconds))
            except pymysql.err.IntegrityError:
                raise HTTPException(409, 'Tên bài thi đã tồn tại trong lớp–môn này. Hãy chọn tên khác.')
            ident = db.lastrowid
            insert_questions(db, items, ident, course_id)
        await changed(teacher_id(request))
        return {'exam_id':ident, 'title':name, 'imported':len(items),
                'time_limit_seconds':time_limit_seconds}

    @router.get('/exams/{exam_id}/editor')
    def exam_editor(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request))
            db.execute('SELECT id,position,text,a,b,c,d,correct_choice,points FROM questions WHERE exam_id=%s ORDER BY position', (exam_id,))
            return {'exam':exam, 'questions':db.fetchall()}

    @router.put('/exams/{exam_id}')
    async def edit_exam(exam_id: int, item: ExamEdit, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request), True)
            if exam['status']!='draft':
                raise HTTPException(409, 'Chỉ sửa nội dung và đáp án của bài thi nháp')
            if exam['revision']!=item.revision:
                raise HTTPException(409, 'Đề đã thay đổi. Hãy mở lại trình chỉnh sửa.')
            db.execute('SELECT 1 FROM answers a JOIN questions q ON q.id=a.question_id WHERE q.exam_id=%s LIMIT 1', (exam_id,))
            if db.fetchone():
                raise HTTPException(409, 'Đề đã có đáp án nộp, không thể sửa nội dung')
            db.execute('DELETE FROM questions WHERE exam_id=%s', (exam_id,))
            insert_questions(db, [q.model_dump() for q in item.questions], exam_id, exam['course_id'])
            try:
                db.execute('UPDATE exams SET title=%s,time_limit_seconds=%s,revision=revision+1 WHERE id=%s',
                           (item.title, item.time_limit_seconds, exam_id))
            except pymysql.err.IntegrityError:
                raise HTTPException(409, 'Tên bài thi đã tồn tại trong lớp–môn này. Hãy chọn tên khác.')
        await changed(teacher_id(request))
        return {'id':exam_id, 'revision':item.revision+1}

    @router.patch('/exams/{exam_id}')
    async def rename_exam(exam_id: int, item: ExamTitle, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request), True)
            if exam['status']=='open' or exam['revision']!=item.revision:
                raise HTTPException(409, 'Bài đang thi hoặc đã thay đổi. Hãy tải lại trước khi đổi tên.')
            try:
                db.execute('UPDATE exams SET title=%s,revision=revision+1 WHERE id=%s', (item.title, exam_id))
            except pymysql.err.IntegrityError:
                raise HTTPException(409, 'Tên bài thi đã tồn tại trong lớp–môn này. Hãy chọn tên khác.')
        await changed(teacher_id(request))
        return {'id':exam_id, 'revision':item.revision+1}

    @router.get('/exams/{exam_id}/check-in')
    def exam_check_in(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request))
            if exam['status'] != 'draft':
                raise HTTPException(409, 'Chỉ chọn thiết bị trước khi bắt đầu bài thi nháp.')
            people = roster(db, exam=exam)
            db.execute('''SELECT d.device_id,d.protocol,(d.token_hash IS NOT NULL) AS wifi_ready,
                (bi.connected=TRUE AND bi.last_seen>=UTC_TIMESTAMP(3)-INTERVAL 15 SECOND) AS ble_online,
                (bi.connected=TRUE AND bi.authorized=TRUE AND bi.last_seen>=UTC_TIMESTAMP(3)-INTERVAL 15 SECOND) AS ble_authorized,
                bi.gatt_profile AS ble_profile
                FROM devices d LEFT JOIN ble_device_inventory bi ON bi.device_id=d.device_id
                WHERE d.teacher_id=%s ORDER BY d.device_id''', (teacher_id(request),))
            devices = db.fetchall()
        return {'exam':{'id':exam_id,'title':exam['title'],'revision':exam['revision']},
                'students':people, 'devices':devices}

    @router.post('/exams/{exam_id}/start')
    async def start_exam(exam_id: int, item: StartExamIn, request: Request):
        try:
            with database() as db:
                exam = own_exam(db, exam_id, teacher_id(request), True)
                if exam['status'] != 'draft':
                    raise HTTPException(409, 'Chỉ bắt đầu bài thi nháp. Dùng Thi lại để giữ kết quả cũ.')
                if exam['revision'] != item.revision:
                    raise HTTPException(409, 'Bài thi đã thay đổi. Hãy chọn lại thiết bị trước khi bắt đầu.')
                own_course(db, exam['course_id'], teacher_id(request), True)
                people = roster(db, exam=exam)
                if not people:
                    raise HTTPException(409, 'Lớp–môn chưa có sinh viên. Hãy thêm danh sách trước khi thi.')
                assignments = validate_exam_assignments(people, item.assignments)
                for assignment in item.assignments:
                    if assignment.protocol is None:
                        continue  # Older API clients did not send a protocol.
                    db.execute('''SELECT protocol,token_hash IS NOT NULL AS wifi_ready
                        FROM devices WHERE teacher_id=%s AND device_id=%s''',
                        (teacher_id(request), assignment.device_id))
                    registered = db.fetchone()
                    if assignment.protocol == 'WIFI' and not (registered and registered['protocol'] == 'WIFI' and registered['wifi_ready']):
                        raise HTTPException(422, 'ESP32 cần chọn mã thiết bị Wi-Fi đã kích hoạt, không nhập token.')
                    if assignment.protocol == 'BLE' and registered and registered['wifi_ready']:
                        raise HTTPException(422, 'Mã này thuộc ESP32 Wi-Fi. Hãy chọn đúng loại thiết bị.')
                db.execute("SELECT id FROM questions WHERE status='open' LIMIT 1")
                if db.fetchone():
                    raise HTTPException(409, 'Gateway hiện chỉ hỗ trợ một lượt thi mở. Hãy kết thúc lượt hiện tại trước.')
                db.execute('SELECT id,position FROM questions WHERE exam_id=%s ORDER BY position LIMIT 1 FOR UPDATE', (exam_id,))
                first = db.fetchone()
                if not first:
                    raise HTTPException(409, 'Bài thi chưa có câu hỏi')
                db.executemany('''INSERT IGNORE INTO devices(teacher_id,device_id) VALUES(%s,%s)''',
                               [(teacher_id(request), device) for device in assignments.values()])
                db.executemany('''INSERT INTO exam_students(exam_id,student_id,mssv,name,device_id)
                    VALUES(%s,%s,%s,%s,%s)''',
                    [(exam_id, person['id'], person['mssv'], person['name'], assignments[person['id']])
                     for person in people])
                db.execute("UPDATE questions SET status='open' WHERE id=%s", (first['id'],))
                db.execute('''UPDATE exams SET status='open',current_position=%s,revision=revision+1,
                    deadline_at=CASE WHEN time_limit_seconds IS NULL THEN NULL
                        ELSE DATE_ADD(UTC_TIMESTAMP(3),INTERVAL time_limit_seconds SECOND) END WHERE id=%s''',
                    (first['position'], exam_id))
                first_position = first['position']
        except pymysql.err.IntegrityError:
            raise HTTPException(409, 'Đang có lượt thi khác mở hoặc mã thiết bị bị trùng trong bài thi.')
        await changed(teacher_id(request))
        return {'id':exam_id, 'status':'open', 'position':first_position}

    @router.post('/exams/{exam_id}/advance')
    async def force_advance(exam_id: int, request: Request,
                            expected_position: int | None = Query(default=None, gt=0)):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request), True)
            if exam['status']!='open':
                raise HTTPException(409, 'Bài thi chưa mở hoặc đã kết thúc')
            if expected_position is not None and exam['current_position'] != expected_position:
                raise HTTPException(409, 'Câu hỏi đã tự chuyển. Hãy xem câu hiện tại trước khi thao tác lại.')
            db.execute("SELECT id FROM questions WHERE exam_id=%s AND status='open' FOR UPDATE", (exam_id,))
            current = db.fetchone()
            if not current:
                raise HTTPException(409, 'Không có câu đang mở')
            result = advance_exam(db, exam_id, current['id'])
        await changed(teacher_id(request))
        return {'id':exam_id, **result}

    @router.post('/exams/{exam_id}/finish')
    async def finish_exam(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request), True)
            if exam['status']!='open':
                raise HTTPException(409, 'Bài thi chưa mở hoặc đã kết thúc')
            db.execute("UPDATE questions SET status='closed' WHERE exam_id=%s", (exam_id,))
            db.execute("UPDATE exams SET status='closed',current_position=NULL,deadline_at=NULL,revision=revision+1 WHERE id=%s", (exam_id,))
        await changed(teacher_id(request))
        return {'id':exam_id, 'status':'closed', 'completed':True, 'position':None}

    @router.post('/exams/{exam_id}/retake', status_code=201)
    async def retake_exam(exam_id: int, request: Request):
        """Copy a finished exam into a new draft; never alter previous answers."""
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request), True)
            if exam['status'] != 'closed':
                raise HTTPException(409, 'Chỉ có thể thi lại bài đã kết thúc')
            db.execute('''SELECT position,text,a,b,c,d,correct_choice,points
                FROM questions WHERE exam_id=%s ORDER BY position''', (exam_id,))
            questions = db.fetchall()
            if not questions:
                raise HTTPException(409, 'Bài thi không còn câu hỏi để thi lại')
            own_course(db, exam['course_id'], teacher_id(request), True)
            title = unique_retake_title(db, exam['course_id'], exam['title'])
            try:
                db.execute('INSERT INTO exams(title,course_id,time_limit_seconds) VALUES(%s,%s,%s)',
                           (title, exam['course_id'], exam['time_limit_seconds']))
            except pymysql.err.IntegrityError:
                raise HTTPException(409, 'Tên lượt thi lại đã thay đổi. Hãy thử lại.')
            new_id = db.lastrowid
            db.executemany('''INSERT INTO questions(position,text,a,b,c,d,correct_choice,points,exam_id,course_id)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
                [(q['position'], q['text'], q['a'], q['b'], q['c'], q['d'], q['correct_choice'],
                  q['points'], new_id, exam['course_id']) for q in questions])
        await changed(teacher_id(request))
        return {'exam_id':new_id, 'source_exam_id':exam_id, 'status':'draft', 'title':title,
                'question_count':len(questions)}

    @router.delete('/exams/{exam_id}')
    async def delete_exam(exam_id: int, request: Request):
        with database() as db:
            own_exam(db, exam_id, teacher_id(request), True)
            purge_exam(db, exam_id)
        await changed(teacher_id(request))
        return {'id':exam_id, 'deleted':True}

    @router.get('/exams/{exam_id}/summary')
    def exam_summary(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request))
            return exam_summary_data(db, exam)

    @router.get('/exams/{exam_id}/scores.xlsx')
    def exam_scores_xlsx(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request))
            if exam['status'] != 'closed':
                raise HTTPException(409, 'Chỉ xuất bảng điểm sau khi bài thi kết thúc.')
            summary = exam_summary_data(db, exam)
        return StreamingResponse(export_score_workbook(summary),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition':f'attachment; filename="bang_diem_bai_thi_{exam_id}.xlsx"',
                     'Cache-Control':'no-store'})

    @router.get('/exams/{exam_id}/presentation')
    def presentation(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db, exam_id, teacher_id(request))
            db.execute('''SELECT CASE WHEN deadline_at IS NULL THEN NULL ELSE
                GREATEST(0,TIMESTAMPDIFF(MICROSECOND,UTC_TIMESTAMP(3),deadline_at) DIV 1000)
                END AS remaining_ms FROM exams WHERE id=%s''', (exam_id,))
            exam['remaining_ms'] = db.fetchone()['remaining_ms']
            course = own_course(db, exam['course_id'], teacher_id(request))
            people = roster(db, exam=exam)
            db.execute('''SELECT q.id,q.position,q.text,q.a,q.b,q.c,q.d,q.points,COUNT(a.id) AS submitted
                FROM questions q LEFT JOIN answers a ON a.question_id=q.id WHERE q.exam_id=%s
                GROUP BY q.id,q.position,q.text,q.a,q.b,q.c,q.d,q.points ORDER BY q.position''', (exam_id,))
            slides = db.fetchall()
        return {'exam':exam,'questions':slides,'total_students':len(people),
                'context':{'class_name':course['class_name'],'subject_name':course['subject_name']}}

    @router.get('/questions')
    def questions(request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request))
            db.execute("""SELECT id,text,a,b,c,d,points,status,
                CASE WHEN status='closed' THEN correct_choice ELSE NULL END AS correct_choice
                FROM questions WHERE course_id=%s AND exam_id IS NULL ORDER BY id DESC""", (course_id,))
            return db.fetchall()

    @router.post('/questions', status_code=201)
    async def create_question(item: QuestionIn, request: Request, course_id: int = Query(gt=0)):
        with database() as db:
            own_course(db, course_id, teacher_id(request), True)
            values = item.model_dump()
            db.execute('INSERT INTO questions(text,a,b,c,d,correct_choice,points,course_id) VALUES(%s,%s,%s,%s,%s,%s,%s,%s)',
                       tuple(values[k] for k in ['text','a','b','c','d','correct_choice','points'])+(course_id,))
            ident = db.lastrowid
        await changed(teacher_id(request))
        return {'id':ident,'status':'draft'}

    @router.delete('/questions/{question_id}')
    async def delete_question(question_id: int, request: Request):
        with database() as db:
            question = own_question(db, question_id, teacher_id(request), True)
            if question['exam_id'] is not None:
                raise HTTPException(409, 'Câu hỏi thuộc bài thi. Hãy xóa bài thi thay vì xóa riêng câu.')
            purge_question(db, question_id)
        await changed(teacher_id(request))
        return {'id':question_id, 'deleted':True}

    @router.post('/questions/{question_id}/open')
    async def open_question(question_id: int, request: Request):
        try:
            with database() as db:
                question = own_question(db, question_id, teacher_id(request), True)
                if question['exam_id'] is not None or question['status']!='draft':
                    raise HTTPException(409, 'Chỉ mở câu hỏi riêng lẻ đang nháp')
                own_course(db, question['course_id'], teacher_id(request), True)
                db.execute('''SELECT COUNT(*) AS total,COUNT(DISTINCT NULLIF(s.device_id,'')) AS devices
                    FROM enrollments en JOIN students s ON s.id=en.student_id
                    WHERE en.course_id=%s AND en.active=TRUE''', (question['course_id'],))
                identities = db.fetchone()
                if identities['total'] != identities['devices']:
                    raise HTTPException(409, 'Câu hỏi riêng lẻ cần mã thiết bị khác nhau cho mọi sinh viên. Với thiết bị dùng chung, hãy tạo bài thi và chọn thiết bị trước khi bắt đầu.')
                db.execute('''INSERT INTO question_students(question_id,student_id,mssv,name,device_id)
                    SELECT %s,s.id,s.mssv,s.name,s.device_id FROM enrollments en JOIN students s ON s.id=en.student_id
                    WHERE en.course_id=%s AND en.active=TRUE''', (question_id,question['course_id']))
                if not db.rowcount:
                    raise HTTPException(409, 'Lớp–môn chưa có sinh viên')
                db.execute("UPDATE questions SET status='open' WHERE id=%s", (question_id,))
        except pymysql.err.IntegrityError:
            raise HTTPException(409, 'Gateway hiện chỉ hỗ trợ một lượt thi mở')
        await changed(teacher_id(request))
        return {'id':question_id,'status':'open'}

    @router.post('/questions/{question_id}/close')
    async def close_question(question_id: int, request: Request):
        with database() as db:
            question = own_question(db, question_id, teacher_id(request), True)
            if question['exam_id'] is not None or question['status']!='open':
                raise HTTPException(409, 'Chỉ đóng câu riêng lẻ đang mở')
            db.execute("UPDATE questions SET status='closed' WHERE id=%s", (question_id,))
        await changed(teacher_id(request))
        return {'id':question_id,'status':'closed'}

    @router.get('/questions/{question_id}/results')
    def results(question_id: int, request: Request):
        with database() as db:
            question = own_question(db, question_id, teacher_id(request))
            if question['exam_id'] is not None:
                raise HTTPException(409, 'Hãy xem tổng kết bài thi cho câu này')
            people = roster(db, question=question)
            db.execute('SELECT student_id,choice,submitted_at FROM answers WHERE question_id=%s', (question_id,))
            answer_map = {a['student_id']:a for a in db.fetchall()}
        for person in people:
            answer = answer_map.get(person['id'], {})
            person.update(choice=answer.get('choice'),submitted_at=answer.get('submitted_at'))
            person['awarded_points'] = rounded_score(question['points']) if question['status']=='closed' and person['choice']==question['correct_choice'] else 0
        if question['status']!='closed':
            question['correct_choice'] = None
        return {'question':question,'students':people,
                'counts':{c:sum(s['choice']==c for s in people) for c in 'ABCD'},
                'submitted':sum(s['choice'] is not None for s in people),'total':len(people)}

    def gateway(key):
        if not key or key!=gateway_key:
            raise HTTPException(401, 'Mã gateway không hợp lệ')

    @router.get('/gateway/current')
    def gateway_current(x_gateway_key: str | None = Header(default=None)):
        gateway(x_gateway_key)
        with database() as db:
            db.execute("SELECT id AS question_id,exam_id,position FROM questions WHERE status='open' AND course_id IS NOT NULL")
            question = db.fetchone()
            if not question:
                raise HTTPException(409, 'Không có câu hỏi đang mở')
            return question

    @router.post('/gateway/authorize')
    def gateway_authorize(item: BleAuthorizeIn, x_gateway_key: str | None = Header(default=None)):
        """Authenticate a known BLE ID immediately; an open quiz also requires roster membership."""
        gateway(x_gateway_key)
        with database() as db:
            return ble_authorization(db, item.device_id)

    @router.post('/gateway/online', status_code=202)
    async def gateway_online(item: BleOnlineIn, x_gateway_key: str | None = Header(default=None)):
        gateway(x_gateway_key)
        with database() as db:
            db.execute('''SELECT connected,gateway_id,gatt_profile,authorized,
                last_seen<UTC_TIMESTAMP(3)-INTERVAL 15 SECOND AS stale
                FROM ble_device_inventory WHERE device_id=%s FOR UPDATE''', (item.device_id,))
            previous = db.fetchone()
            changed_state = (previous is None or bool(previous['connected']) != item.connected
                             or (item.connected and bool(previous['stale']))
                             or (item.connected and previous['gateway_id'] != item.gateway_id)
                             or (item.connected and bool(previous['authorized']) != item.authorized)
                             or (item.connected and item.gatt_profile is not None
                                 and previous['gatt_profile'] != item.gatt_profile))
            db.execute('''INSERT INTO ble_device_inventory(device_id,gateway_id,connected,last_seen,gatt_profile,authorized)
                VALUES(%s,%s,%s,UTC_TIMESTAMP(3),%s,%s) ON DUPLICATE KEY UPDATE
                gateway_id=VALUES(gateway_id),connected=VALUES(connected),last_seen=VALUES(last_seen),
                gatt_profile=COALESCE(VALUES(gatt_profile),gatt_profile),authorized=VALUES(authorized)''',
                (item.device_id, item.gateway_id, item.connected,
                 item.gatt_profile, item.authorized if item.connected else False))
            owners = []
            if changed_state:
                db.execute('SELECT teacher_id FROM devices WHERE device_id=%s', (item.device_id,))
                owners = [row['teacher_id'] for row in db.fetchall() if row['teacher_id'] is not None]
        for owner in owners:
            await changed(owner)
        return {'accepted':True}

    @router.post('/answers', status_code=201)
    async def submit_answer(item: AnswerIn, x_gateway_key: str | None = Header(default=None)):
        gateway(x_gateway_key)
        with database() as db:
            db.execute('SELECT exam_id FROM questions WHERE id=%s', (item.question_id,))
            peek = db.fetchone()
            if not peek:
                raise HTTPException(409, 'Câu hỏi không còn tồn tại')
            if peek['exam_id'] is not None:
                db.execute('''SELECT status,(deadline_at IS NOT NULL AND deadline_at<=UTC_TIMESTAMP(3)) AS expired
                    FROM exams WHERE id=%s FOR UPDATE''', (peek['exam_id'],))
                exam = db.fetchone()
                if not exam or exam['status']!='open':
                    raise HTTPException(409, 'Bài thi đã đóng')
                if exam['expired']:
                    raise HTTPException(409, 'Đã hết giờ câu hỏi. Không nhận đáp án trễ.')
            db.execute("SELECT * FROM questions WHERE id=%s AND status='open' AND course_id IS NOT NULL FOR UPDATE", (item.question_id,))
            question = db.fetchone()
            if not question:
                raise HTTPException(409, 'Câu hỏi đã đóng hoặc chưa mở. Không nhận phiếu của câu cũ.')
            if question['exam_id'] is not None:
                db.execute('SELECT student_id FROM exam_students WHERE exam_id=%s AND device_id=%s', (question['exam_id'],item.device_id.strip()))
            else:
                db.execute('SELECT student_id FROM question_students WHERE question_id=%s AND device_id=%s', (question['id'],item.device_id.strip()))
            student = db.fetchone()
            if not student:
                raise HTTPException(403, 'Thiết bị không thuộc danh sách tham gia lượt thi này')
            try:
                db.execute('INSERT INTO answers(question_id,student_id,choice,sequence) VALUES(%s,%s,%s,%s)',
                           (question['id'],student['student_id'],item.choice,item.sequence))
            except pymysql.err.IntegrityError:
                raise HTTPException(409, 'Sinh viên đã nộp câu hỏi này')
            progress = {'advanced':False,'completed':False,'position':question['position']}
            if question['exam_id'] is not None:
                db.execute('SELECT COUNT(*) AS count FROM exam_students WHERE exam_id=%s', (question['exam_id'],))
                total = db.fetchone()['count']
                db.execute('SELECT COUNT(*) AS count FROM answers WHERE question_id=%s', (question['id'],))
                if total and db.fetchone()['count']>=total:
                    progress = advance_exam(db,question['exam_id'],question['id'])
            db.execute('SELECT c.teacher_id FROM class_subjects cs JOIN classes c ON c.id=cs.class_id WHERE cs.id=%s', (question['course_id'],))
            owner = db.fetchone()['teacher_id']
        await changed(owner)
        return {'question_id':question['id'],'student_id':student['student_id'],'choice':item.choice,'accepted':True,**progress}

    @router.post('/gateway/presence', status_code=202)
    async def device_presence(item: PresenceIn, x_gateway_key: str | None = Header(default=None)):
        gateway(x_gateway_key)
        with database() as db:
            db.execute("SELECT id,exam_id,course_id FROM questions WHERE id=%s AND status='open'", (item.question_id,))
            question = db.fetchone()
            if not question:
                raise HTTPException(409, 'Câu hỏi đã đóng hoặc chưa mở. Không nhận tín hiệu của lượt thi cũ.')
            if question['exam_id'] is not None:
                db.execute('''SELECT status,(deadline_at IS NOT NULL AND deadline_at<=UTC_TIMESTAMP(3)) AS expired
                    FROM exams WHERE id=%s''', (question['exam_id'],))
                exam = db.fetchone()
                if not exam or exam['status']!='open':
                    raise HTTPException(409, 'Bài thi đã đóng')
                if exam['expired']:
                    raise HTTPException(409, 'Đã hết giờ câu hỏi. Không nhận tín hiệu của câu cũ.')
                db.execute('SELECT student_id FROM exam_students WHERE exam_id=%s AND device_id=%s LIMIT 2',
                           (question['exam_id'],item.device_id.strip()))
            else:
                db.execute('SELECT student_id FROM question_students WHERE question_id=%s AND device_id=%s LIMIT 2',
                           (question['id'],item.device_id.strip()))
            matches = db.fetchall()
            if not matches:
                raise HTTPException(403, 'Thiết bị không thuộc danh sách tham gia lượt thi này')
            if len(matches)>1:
                raise HTTPException(409, 'Mã thiết bị bị trùng trong cùng lượt thi')
            student_id = matches[0]['student_id']
            db.execute('''SELECT c.teacher_id FROM class_subjects cs JOIN classes c ON c.id=cs.class_id
                WHERE cs.id=%s''', (question['course_id'],))
            owner = db.fetchone()['teacher_id']
            db.execute('''INSERT INTO device_presence(student_id,exam_id,protocol,gateway_id,connected,rssi,last_seen)
                VALUES(%s,%s,%s,%s,%s,%s,UTC_TIMESTAMP(3)) ON DUPLICATE KEY UPDATE
                exam_id=VALUES(exam_id),protocol=VALUES(protocol),gateway_id=VALUES(gateway_id),
                connected=VALUES(connected),rssi=VALUES(rssi),last_seen=VALUES(last_seen)''',
                (student_id,question['exam_id'],item.protocol,item.gateway_id,item.connected,item.rssi))
            db.execute('''INSERT INTO device_reports(student_id,exam_id,protocol,connected,rssi,received_at)
                VALUES(%s,%s,%s,%s,%s,UTC_TIMESTAMP(3))''',
                (student_id,question['exam_id'],item.protocol,item.connected,item.rssi))
            db.execute('DELETE FROM device_reports WHERE student_id=%s AND received_at<UTC_TIMESTAMP(3)-INTERVAL 60 SECOND', (student_id,))
        await changed(owner)
        return {'accepted':True}

    def wifi_student(token):
        if not token or len(token)>128:
            raise HTTPException(401, 'Mã thiết bị Wi-Fi không hợp lệ')
        token_hash = hashlib.sha256(token.encode()).hexdigest()
        with database() as db:
            db.execute('''SELECT teacher_id,device_id FROM devices WHERE token_hash=%s AND protocol='WIFI' ''',
                (token_hash,))
            device = db.fetchone()
        if not device:
            raise HTTPException(401, 'Mã thiết bị Wi-Fi không hợp lệ')
        return device

    def wifi_question(db, question_id, device):
        db.execute('''SELECT q.id AS question_id,q.exam_id,q.position,c.teacher_id
            FROM questions q JOIN class_subjects cs ON cs.id=q.course_id
            JOIN classes c ON c.id=cs.class_id
            WHERE q.id=%s AND q.status='open' ''', (question_id,))
        question = db.fetchone()
        if not question:
            raise HTTPException(409, 'Không có câu hỏi này đang mở')
        if question['teacher_id'] != device['teacher_id']:
            raise HTTPException(403, 'Thiết bị không thuộc giảng viên của lượt thi này')
        if question['exam_id'] is not None:
            db.execute('SELECT student_id FROM exam_students WHERE exam_id=%s AND device_id=%s',
                       (question['exam_id'],device['device_id']))
        else:
            db.execute('SELECT student_id FROM question_students WHERE question_id=%s AND device_id=%s',
                       (question_id,device['device_id']))
        if not db.fetchone():
            raise HTTPException(403, 'Thiết bị không tham gia câu hỏi này')
        if question['exam_id'] is not None:
            db.execute('''SELECT status,(deadline_at IS NOT NULL AND deadline_at<=UTC_TIMESTAMP(3)) AS expired
                FROM exams WHERE id=%s''', (question['exam_id'],))
            exam = db.fetchone()
            if not exam or exam['status']!='open' or exam['expired']:
                raise HTTPException(409, 'Bài thi hoặc câu hỏi đã hết giờ')
        return question

    @router.get('/device/ping')
    def wifi_ping(x_device_token: str | None = Header(default=None)):
        device = wifi_student(x_device_token)
        return {'connected': True, 'device_id': device['device_id']}

    @router.get('/device/current')
    def wifi_current(x_device_token: str | None = Header(default=None)):
        device = wifi_student(x_device_token)
        with database() as db:
            db.execute("SELECT id FROM questions WHERE status='open' AND course_id IS NOT NULL LIMIT 1")
            current = db.fetchone()
            if not current:
                raise HTTPException(409, 'Chưa có câu hỏi đang mở')
            question = wifi_question(db, current['id'], device)
            return {key:question[key] for key in ('question_id','exam_id','position')}

    @router.post('/device/answers', status_code=201)
    async def wifi_answer(item: WifiAnswerIn, x_device_token: str | None = Header(default=None)):
        device = wifi_student(x_device_token)
        with database() as db:
            wifi_question(db, item.question_id, device)
        result = await submit_answer(AnswerIn(question_id=item.question_id,device_id=device['device_id'],
                                              choice=item.choice,sequence=item.sequence), gateway_key)
        return {**result, 'ack': 'ACK'}

    @router.post('/device/presence', status_code=202)
    async def wifi_presence(item: WifiPresenceIn, x_device_token: str | None = Header(default=None)):
        device = wifi_student(x_device_token)
        with database() as db:
            wifi_question(db, item.question_id, device)
        return await device_presence(PresenceIn(question_id=item.question_id,device_id=device['device_id'],
                                               protocol='WIFI',gateway_id=item.gateway_id,
                                               connected=item.connected,rssi=item.rssi), gateway_key)

    @router.get('/exams/{exam_id}/devices')
    def exam_devices(exam_id: int, request: Request):
        with database() as db:
            exam = own_exam(db,exam_id,teacher_id(request))
            people = roster(db,exam=exam)
            db.execute("SELECT device_id FROM devices WHERE teacher_id=%s AND protocol='WIFI' AND token_hash IS NOT NULL",
                       (teacher_id(request),))
            wifi_devices = {row['device_id'] for row in db.fetchall()}
            db.execute('''SELECT p.student_id,p.protocol,p.gateway_id,p.connected,p.rssi,p.last_seen,
                TIMESTAMPDIFF(SECOND,p.last_seen,UTC_TIMESTAMP()) AS age_seconds,
                (SELECT COUNT(*) FROM device_reports r WHERE r.student_id=p.student_id AND r.exam_id=p.exam_id
                    AND r.received_at>=UTC_TIMESTAMP()-INTERVAL 60 SECOND AND r.connected=TRUE) AS recent_reports,
                (SELECT COUNT(*) FROM device_reports r WHERE r.student_id=p.student_id AND r.exam_id=p.exam_id
                    AND r.received_at>=UTC_TIMESTAMP()-INTERVAL 60 SECOND AND r.connected=FALSE) AS recent_disconnects
                FROM device_presence p JOIN students s ON s.id=p.student_id
                WHERE s.teacher_id=%s AND p.exam_id=%s''', (teacher_id(request),exam_id))
            presence = {row['student_id']:row for row in db.fetchall()}
            inventory = {}
            if exam['status'] in ('draft', 'open') and people:
                placeholders = ','.join(['%s'] * len(people))
                db.execute(f'''SELECT device_id,gateway_id,connected,gatt_profile,authorized,
                    TIMESTAMPDIFF(SECOND,last_seen,UTC_TIMESTAMP()) AS age_seconds
                    FROM ble_device_inventory WHERE device_id IN ({placeholders})''',
                    tuple(student['device_id'] for student in people))
                inventory = {row['device_id']:row for row in db.fetchall()}
            db.execute("SELECT id FROM questions WHERE exam_id=%s AND status='open'", (exam_id,))
            current = db.fetchone()
            submitted = set()
            if current:
                db.execute('SELECT student_id FROM answers WHERE question_id=%s', (current['id'],))
                submitted = {row['student_id'] for row in db.fetchall()}
        devices = []
        for student in people:
            row = presence.get(student['id'])
            ble_inventory = inventory.get(student['device_id'])
            if row is None and student['device_id'] in inventory:
                row = {**ble_inventory, 'protocol':'BLE', 'rssi':None,
                       'recent_reports':0, 'recent_disconnects':0}
            status = 'unknown'
            stability = 'unknown'
            if row:
                age = max(0,row['age_seconds'] or 0)
                if row['protocol']=='SIMULATED':
                    status = 'simulated' if row['connected'] and age<=15 else 'disconnected' if not row['connected'] else 'stale'
                    stability = 'simulated'
                elif not row['connected']:
                    status = 'disconnected'
                elif age>15:
                    status = 'stale'
                else:
                    status = 'connected'
                    stability = 'weak' if row['rssi'] is not None and row['rssi'] < -85 else 'stable' if row['recent_reports']>=3 and not row['recent_disconnects'] and (row['protocol']!='BLE' or row['rssi'] is not None) else 'observing'
            devices.append({**student,'status':status,'stability':stability,
                            'expected_protocol':'WIFI' if student['device_id'] in wifi_devices else 'BLE',
                            'protocol':row['protocol'] if row else None,
                            'ble_profile':ble_inventory['gatt_profile'] if ble_inventory else None,
                            'ble_authorized':bool(ble_inventory['authorized']) if ble_inventory else False,
                            'gateway_id':row['gateway_id'] if row else None,
                            'rssi':row['rssi'] if row else None,
                            'age_seconds':max(0,row['age_seconds']) if row and row['age_seconds'] is not None else None,
                            'recent_reports':row['recent_reports'] if row else 0,
                            'submitted_current':student['id'] in submitted})
        return {'exam':{'id':exam_id,'title':exam['title'],'status':exam['status']},
                'total':len(devices),'connected':sum(d['status']=='connected' for d in devices),
                'simulated':sum(d['status']=='simulated' for d in devices),'devices':devices}

    app.include_router(router)
