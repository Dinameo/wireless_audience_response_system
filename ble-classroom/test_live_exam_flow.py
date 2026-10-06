"""Optional HTTP/MySQL integration test. Run only against a fresh disposable DB."""
import http.cookiejar
import json
import os
import unittest
from io import BytesIO
from urllib.error import HTTPError
from urllib.request import HTTPCookieProcessor, Request, build_opener

from openpyxl import load_workbook


@unittest.skipUnless(os.getenv('EXAM_TEST_DB') and os.getenv('DB_NAME') == os.getenv('EXAM_TEST_DB'),
                     'Set EXAM_TEST_DB to the name of a fresh disposable database')
class LiveExamFlow(unittest.TestCase):
    def test_reuse_wifi_device_and_export(self):
        from app import database

        base = os.getenv('EXAM_TEST_URL', 'http://127.0.0.1:8012')
        browser = build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))

        def call(path, method='GET', body=None, headers=None, raw=False):
            payload = None if body is None else json.dumps(body).encode()
            request = Request(base + '/api' + path, data=payload, method=method,
                              headers={**({'Content-Type':'application/json'} if body is not None else {}),
                                       **(headers or {})})
            try:
                with browser.open(request, timeout=10) as response:
                    result = response.read()
                    return result if raw else json.loads(result)
            except HTTPError as error:
                return error.code, json.loads(error.read())

        call('/auth/register', 'POST',
             {'name':'Giảng viên test','email':'integration@example.test',
              'password':'test-password-123'})
        subject = call('/subjects', 'POST', {'code':'IOT','name':'IoT'})
        course1 = call('/classes', 'POST', {'code':'L1','name':'Lớp 1',
                                          'subject_id':subject['id']})
        course2 = call('/classes', 'POST', {'code':'L2','name':'Lớp 2',
                                          'subject_id':subject['id']})
        first = call(f"/classes/{course1['id']}/students", 'POST',
                     {'mssv':'001','name':'Sinh viên một','device_id':'ESP-01',
                      'preferred_protocol':'WIFI'})
        second = call(f"/classes/{course1['id']}/students", 'POST',
                      {'mssv':'002','name':'Sinh viên hai','device_id':''})
        third = call(f"/classes/{course2['id']}/students", 'POST',
                     {'mssv':'003','name':'Sinh viên ba','device_id':'ESP-01',
                      'preferred_protocol':'WIFI'})
        self.assertNotEqual(first['id'], third['id'])
        issued = call(f"/students/{first['id']}/wifi-token?course_id={course1['course_id']}", 'POST')

        def make_exam(course, title):
            with database() as db:
                db.execute('INSERT INTO exams(title,course_id) VALUES(%s,%s)', (title,course))
                exam_id = db.lastrowid
                db.execute('''INSERT INTO questions(text,a,b,c,d,correct_choice,points,exam_id,position,course_id)
                    VALUES('Câu 1','A','B','C','D','A',5,%s,1,%s),
                          ('Câu 2','A','B','C','D','B',5,%s,2,%s)''',
                    (exam_id,course,exam_id,course))
            return exam_id

        exam1 = make_exam(course1['course_id'], 'Kiểm tra 1')
        checked = call(f'/exams/{exam1}/check-in')
        self.assertEqual(len(checked['students']), 2)
        self.assertEqual({student['mssv']:student['preferred_protocol'] for student in checked['students']},
                         {'001':'WIFI','002':'BLE'})
        self.assertIn('ESP-01', {device['device_id'] for device in checked['devices']})
        duplicate = {'revision':checked['exam']['revision'], 'assignments':[
            {'student_id':first['id'],'device_id':'ESP-01'},
            {'student_id':second['id'],'device_id':'ESP-01'}]}
        self.assertEqual(call(f'/exams/{exam1}/start', 'POST', duplicate)[0], 422)
        wrong_wifi = {'revision':checked['exam']['revision'], 'assignments':[
            {'student_id':first['id'],'device_id':'NRF-01','protocol':'BLE'},
            {'student_id':second['id'],'device_id':'not-an-esp32','protocol':'WIFI'}]}
        self.assertEqual(call(f'/exams/{exam1}/start', 'POST', wrong_wifi)[0], 422)
        selected = {'revision':checked['exam']['revision'], 'assignments':[
            {'student_id':first['id'],'device_id':'NRF-01','protocol':'BLE'},
            {'student_id':second['id'],'device_id':'ESP-01','protocol':'WIFI'}]}
        self.assertEqual(call(f'/exams/{exam1}/start', 'POST', selected)['status'], 'open')
        token_header = {'X-Device-Token':issued['token']}
        self.assertEqual(call('/device/ping', headers=token_header),
                         {'connected': True, 'device_id': 'ESP-01'})
        current = call('/device/current', headers=token_header)
        submitted = call('/device/answers', 'POST',
                         {'question_id':current['question_id'],'choice':'A','sequence':1},
                         token_header)
        self.assertEqual(submitted['student_id'], second['id'])
        self.assertEqual(submitted['ack'], 'ACK')
        call(f'/exams/{exam1}/finish', 'POST')
        summary = call(f'/exams/{exam1}/summary')
        self.assertEqual({person['id']:person['score'] for person in summary['students']},
                         {first['id']:0, second['id']:5})
        workbook = load_workbook(BytesIO(call(f'/exams/{exam1}/scores.xlsx', raw=True)))
        self.assertEqual(workbook['Bảng điểm'].max_row, 3)
        self.assertEqual(workbook['Chi tiết câu trả lời'].max_row, 5)

        exam2 = make_exam(course2['course_id'], 'Kiểm tra 2')
        checked2 = call(f'/exams/{exam2}/check-in')
        call(f'/exams/{exam2}/start', 'POST', {'revision':checked2['exam']['revision'],
             'assignments':[{'student_id':third['id'],'device_id':'ESP-01','protocol':'WIFI'}]})
        current2 = call('/device/current', headers=token_header)
        submitted2 = call('/device/answers', 'POST',
                          {'question_id':current2['question_id'],'choice':'A','sequence':2},
                          token_header)
        self.assertEqual(submitted2['student_id'], third['id'])


if __name__ == '__main__':
    unittest.main()
