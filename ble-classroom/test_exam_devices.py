"""Offline checks for reusable per-exam devices and the Excel score export."""
import unittest
import hashlib
from contextlib import contextmanager
from io import BytesIO

from fastapi import HTTPException
from fastapi import FastAPI
from openpyxl import Workbook, load_workbook

from exam_api import DeviceAssignment, configure_exam_api, export_score_workbook, validate_exam_assignments
from importers import parse_students


class ExamDeviceTests(unittest.TestCase):
    def setUp(self):
        self.students = [{'id': 1}, {'id': 2}]

    def test_same_device_can_be_reused_in_another_exam(self):
        picked = [DeviceAssignment(student_id=1, device_id='NRF-01'),
                  DeviceAssignment(student_id=2, device_id='ESP-02')]
        self.assertEqual(validate_exam_assignments(self.students, picked),
                         {1: 'NRF-01', 2: 'ESP-02'})
        self.assertEqual(validate_exam_assignments([{'id': 3}],
                         [DeviceAssignment(student_id=3, device_id='NRF-01')]), {3: 'NRF-01'})

    def test_duplicate_or_missing_device_is_rejected(self):
        with self.assertRaises(HTTPException):
            validate_exam_assignments(self.students,
                [DeviceAssignment(student_id=1, device_id='NRF-01'),
                 DeviceAssignment(student_id=2, device_id='NRF-01')])
        with self.assertRaises(HTTPException):
            validate_exam_assignments(self.students,
                [DeviceAssignment(student_id=1, device_id='NRF-01'),
                 DeviceAssignment(student_id=2, device_id='nrf-01')])
        with self.assertRaises(HTTPException):
            validate_exam_assignments(self.students,
                [DeviceAssignment(student_id=1, device_id='NRF-01')])

    def test_excel_import_needs_only_student_identity(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['MSSV', 'Họ tên'])
        sheet.append(['00123', 'Sinh viên A'])
        output = BytesIO()
        book.save(output)
        self.assertEqual(parse_students(output.getvalue(), 'students.xlsx'),
                         [{'mssv':'00123','name':'Sinh viên A','device_id':''}])

    def test_excel_export_has_scores_answers_and_no_formula(self):
        data = {'max_score':10.0, 'students':[
            {'mssv':'=1+1', 'name':'Sinh viên A', 'device_id':'NRF-01', 'score':5.0,
             'answers':[{'position':1,'choice':'A','correct_choice':'A','correct':True,'points':5.0},
                        {'position':2,'choice':None,'correct_choice':'B','correct':False,'points':0.0}]}
        ]}
        book = load_workbook(export_score_workbook(data))
        scores, detail = book['Bảng điểm'], book['Chi tiết câu trả lời']
        self.assertEqual(scores['G2'].value, 5.0)
        self.assertEqual(scores['D2'].value, 1)
        self.assertEqual(detail['G3'].value, 'Chưa trả lời')
        self.assertEqual(scores['A2'].data_type, 's')
        self.assertTrue(scores['A2'].value.startswith("'"))

    def test_wifi_ping_requires_a_valid_device_token(self):
        token_hash = hashlib.sha256(b'test-device-token').hexdigest()

        class DeviceLookup:
            def __init__(self):
                self.row = None

            def execute(self, query, params):
                self.row = ({'teacher_id': 1, 'device_id': 'ESP-01'}
                            if params[0] == token_hash else None)

            def fetchone(self):
                return self.row

        @contextmanager
        def database():
            yield DeviceLookup()

        async def changed(_owner):
            pass

        app = FastAPI()
        configure_exam_api(app, database, changed, 'test-gateway-key')
        included = next(route.original_router for route in app.routes
                        if hasattr(route, 'original_router'))
        ping = next(route.endpoint for route in included.routes if route.path == '/api/device/ping')
        with self.assertRaises(HTTPException) as missing:
            ping(None)
        self.assertEqual(missing.exception.status_code, 401)
        with self.assertRaises(HTTPException) as wrong:
            ping('wrong')
        self.assertEqual(wrong.exception.status_code, 401)
        self.assertEqual(ping('test-device-token'), {'connected': True, 'device_id': 'ESP-01'})


if __name__ == '__main__':
    unittest.main()
